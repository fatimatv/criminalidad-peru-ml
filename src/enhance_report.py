from __future__ import annotations

import json
from pathlib import Path
from textwrap import dedent

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "outputs" / "tables"
FIGURES = ROOT / "outputs" / "figures"
REPORT = ROOT / "report"


def load_inputs() -> dict:
    return {
        "summary": json.loads((TABLES / "analysis_summary.json").read_text(encoding="utf-8")),
        "merge": json.loads((TABLES / "merge_audit.json").read_text(encoding="utf-8")),
        "clusters": pd.read_csv(TABLES / "cluster_profiles.csv"),
        "clf": pd.read_csv(TABLES / "classification_benchmark.csv"),
        "importance": pd.read_csv(TABLES / "classification_permutation_importance.csv"),
        "clustering": pd.read_csv(TABLES / "clustering_benchmark.csv"),
    }


def fmt(value: float, digits: int = 2) -> str:
    if pd.isna(value):
        return "n.d."
    return f"{value:,.{digits}f}"


def pct(value: float) -> str:
    return f"{fmt(value, 1)}%"


def paragraphs() -> list[tuple[str, str]]:
    data = load_inputs()
    summary = data["summary"]
    merge = data["merge"]
    by_year = pd.DataFrame(summary["eda"]["by_year"])
    clusters = data["clusters"]
    clf = data["clf"]
    importance = data["importance"]
    rf = clf.loc[clf["modelo"].eq("RandomForest")].iloc[0]
    dummy = clf.loc[clf["modelo"].eq("DummyClassifier")].iloc[0]
    top_features = ", ".join(
        f"{row.feature.replace('_', ' ')} ({row.importance_mean:.3f})"
        for row in importance.head(4).itertuples()
    )
    y2020 = by_year.loc[by_year.anio == 2020].iloc[0]
    y2023 = by_year.loc[by_year.anio == 2023].iloc[0]
    y2024 = by_year.loc[by_year.anio == 2024].iloc[0]
    cluster_text = " ".join(
        f"Cluster {int(row.cluster)}: {int(row.observaciones)} observaciones, tasa media {fmt(row.denuncias_tasa_100k,1)}, "
        f"victimizacion {pct(row.victimization)}, percepcion {pct(row.percepcion)} y confianza {pct(row.confianza)}."
        for row in clusters.itertuples()
    )

    return [
        (
            "PORTADA",
            "Seguridad ciudadana en el Peru: identificacion de perfiles territoriales de criminalidad y clasificacion de contextos de alta victimizacion mediante Machine Learning\n\n"
            "Curso: Machine Learning\nEquipo: [Completar integrantes]\nInstitucion: [Completar]\nFecha: septiembre de 2026",
        ),
        (
            "RESUMEN EJECUTIVO",
            f"Este informe presenta un pipeline reproducible de Machine Learning aplicado a seguridad ciudadana en el Peru. "
            f"Se integraron {merge['denuncias_rows_raw']:,} registros de denuncias policiales, {merge['enapres_rows_raw']:,} filas ENAPRES y poblacion oficial proyectada. "
            f"La unidad de analisis viable fue departamento-anio, con {merge['integrated_rows_2018_2024']} observaciones entre {min(merge['integrated_years'])} y {max(merge['integrated_years'])}. "
            "El trabajo no busca predecir delitos individuales: busca describir patrones territoriales agregados, clasificar contextos de alta victimizacion relativa y discutir sus limites tecnicos, juridicos y eticos.",
        ),
        (
            "A. CONTEXTO DEL PROBLEMA",
            "La seguridad ciudadana combina hechos registrados, experiencias de victimizacion, percepcion social de inseguridad y confianza institucional. "
            "Estos cuatro elementos no son equivalentes: una denuncia policial depende de que el hecho ocurra, sea reconocido como delito, sea denunciado y sea registrado; la victimizacion ENAPRES mide experiencia reportada en encuesta; la percepcion resume expectativas y temores; y la confianza refleja relacion con la PNP. "
            "Por ello, el valor academico del proyecto esta en analizar esas dimensiones de manera integrada, evitando conclusiones simplistas sobre criminalidad real.",
        ),
        (
            "B. COMPRENSION DEL NEGOCIO",
            "La institucion usuaria hipotetica seria el Ministerio del Interior o un observatorio de seguridad ciudadana. La necesidad publica es contar con evidencia territorial para priorizar diagnosticos, comparar contextos y comunicar hallazgos de forma transparente. "
            "El producto esperado es un tablero explicativo y un modelo reproducible que apoyen la discusion tecnica. No se propone un sistema de vigilancia, scoring de personas ni asignacion automatica de recursos policiales.",
        ),
        (
            "C. COMPRENSION Y VALIDACION DE DATOS",
            f"La integracion se limito al nivel {merge['merge_level']} porque esa es la granularidad comun validada entre las fuentes. "
            "Aunque SIDPOL contiene informacion mas detallada por fecha, distrito y modalidad, ENAPRES esta publicada para este recurso en agregados nacional/departamental/anual. "
            f"Forzar una union distrital habria multiplicado registros y creado precision falsa. El panel final cubre {len(merge['integrated_departments'])} jurisdicciones y no presenta duplicados departamento-anio.",
        ),
        (
            "D. ANALISIS EXPLORATORIO",
            f"Las denuncias totales agregadas disminuyeron en 2020 hasta {int(y2020.denuncias_total):,}, un comportamiento consistente con un periodo excepcional de movilidad y registro. "
            f"Luego aumentaron hasta {int(y2023.denuncias_total):,} en 2023 y se mantuvieron altas en 2024 con {int(y2024.denuncias_total):,}. "
            f"En 2024, la victimizacion promedio departamental fue {pct(y2024.victimizacion_pct)}, la percepcion de inseguridad {pct(y2024.percepcion_inseguridad_pct)} y la confianza en la PNP {pct(y2024.confianza_pnp_pct)}. "
            "El EDA muestra que la percepcion puede permanecer elevada incluso cuando la victimizacion no crece en la misma magnitud, por eso se construyo la brecha percepcion-victimizacion como indicador interpretativo.",
        ),
        (
            "E. PLANTEAMIENTO DE MACHINE LEARNING",
            "Se formularon dos tareas. La primera es no supervisada: identificar perfiles territoriales mediante clustering con variables de denuncias, victimizacion, percepcion y confianza. "
            "La segunda es supervisada: clasificar departamento-anio con alta victimizacion relativa, definida como estar en el percentil 75 de su anio. "
            "La definicion anual evita que un solo cambio temporal global distorsione el target y permite comparar territorios dentro de cada periodo.",
        ),
        (
            "F. IMPLEMENTACION DE MODELOS",
            "Para clustering se compararon K-Means y clustering aglomerativo Ward con k entre 2 y 7, usando variables estandarizadas. "
            "Para clasificacion se compararon DummyClassifier, Regresion Logistica y Random Forest. La particion fue temporal: entrenamiento en 2019-2022 y prueba en 2023-2024. "
            "Esta decision es mas exigente que una division aleatoria porque simula el uso real del modelo sobre anios futuros.",
        ),
        (
            "G. RESULTADOS DEL CLUSTERING",
            f"La seleccion final fue {summary['clustering']['best_model']} con k={summary['clustering']['k']}. "
            "La separacion no es perfecta: el silhouette es moderado, lo cual es esperable en datos sociales agregados donde los territorios no se dividen en grupos totalmente puros. "
            f"Aun asi, los perfiles son interpretables. {cluster_text} "
            "Estos grupos deben leerse como perfiles analiticos, no como etiquetas permanentes ni como rankings de peligrosidad.",
        ),
        (
            "H. RESULTADOS DE CLASIFICACION",
            f"El baseline DummyClassifier obtuvo accuracy {dummy.accuracy:.3f}, pero precision, recall y F1 iguales a cero para la clase de alta victimizacion; por tanto, la accuracy aislada habria sido enganosa. "
            f"Random Forest fue el mejor modelo con accuracy {rf.accuracy:.3f}, precision {rf.precision:.3f}, recall {rf.recall:.3f}, F1 {rf.f1:.3f}, ROC-AUC {rf.roc_auc:.3f} y PR-AUC {rf.pr_auc:.3f}. "
            "El recall alto es importante porque reduce falsos negativos, es decir, casos de alta victimizacion que el sistema no detectaria. Sin embargo, la muestra de prueba es pequena, por lo que las metricas deben actualizarse y validarse con nuevos anios.",
        ),
        (
            "I. EXPLICABILIDAD",
            f"La importancia por permutacion indico que las variables con mayor aporte predictivo fueron: {top_features}. "
            "Estas importancias no prueban causalidad. Indican que, dentro de este experimento, al alterar esas variables se deteriora el desempeno del modelo. "
            "La lectura correcta es tecnica y prudente: los rezagos de denuncias y percepcion ayudan a anticipar contextos de victimizacion alta relativa, pero no explican por si solos por que ocurre la inseguridad.",
        ),
        (
            "J. PERSPECTIVA JURIDICA Y DE POLITICA PUBLICA",
            "Un modelo aplicado al Estado debe respetar legalidad, transparencia, proporcionalidad, calidad de datos, rendicion de cuentas y supervision humana. "
            "En el marco peruano son relevantes la Ley de Gobierno Digital, la Ley de Proteccion de Datos Personales y la Politica Nacional Multisectorial de Seguridad Ciudadana al 2030. "
            "Aunque el panel usado es agregado y no contiene datos personales directos, sus resultados pueden afectar narrativas publicas sobre territorios; por eso se recomienda evitar lenguaje estigmatizante y documentar siempre las limitaciones.",
        ),
        (
            "K. INTELIGENCIA ARTIFICIAL RESPONSABLE",
            "Los principales riesgos son sesgo de reporte, subregistro, diferencias territoriales en acceso a comisarias, cambios administrativos, sesgo de encuesta, ecological fallacy y automation bias. "
            "Las mitigaciones propuestas son: publicar diccionario de variables, mantener trazabilidad de fuentes, recalcular metricas con cada actualizacion, incluir revision experta, monitorear desempeno por periodos, prohibir usos individuales y presentar el tablero como apoyo al analisis, no como veredicto automatizado.",
        ),
        (
            "CONCLUSIONES",
            "El proyecto demuestra que es posible integrar datos oficiales para construir un diagnostico territorial reproducible de seguridad ciudadana en el Peru. "
            "La decision metodologica mas importante fue trabajar en departamento-anio para respetar la estructura real de las fuentes. "
            "El clustering aporta perfiles utiles para exploracion y la clasificacion muestra senales predictivas superiores al baseline. "
            "La principal advertencia es que las denuncias no son criminalidad real completa, y las salidas del modelo deben usarse siempre con interpretacion humana y contexto institucional.",
        ),
        (
            "RECOMENDACIONES",
            "Actualizar el pipeline cuando MININTER e INEI publiquen nuevos cortes, incorporar covariables oficiales solo si comparten granularidad, revisar el umbral de alta victimizacion con especialistas, agregar intervalos de incertidumbre cuando sea posible y complementar el tablero con lectura cualitativa territorial. "
            "Para exposicion, se recomienda enfatizar que el proyecto combina rigor tecnico con prudencia publica: modela patrones agregados, no personas.",
        ),
        (
            "REFERENCIAS IEEE",
            (ROOT / "references" / "referencias_ieee.md").read_text(encoding="utf-8"),
        ),
    ]


def style_docx_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt_p = p.paragraph_format
    fmt_p.first_line_indent = Cm(1.27)
    fmt_p.line_spacing = 2.0
    fmt_p.space_before = Pt(0)
    fmt_p.space_after = Pt(0)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(12)


def add_docx_table(doc: Document, df: pd.DataFrame, title: str):
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = fmt(value, 3) if isinstance(value, float) else str(value)


def write_docx() -> None:
    data = load_inputs()
    REPORT.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(12)

    for title, body in paragraphs():
        doc.add_heading(title, level=1)
        for raw in body.split("\n\n"):
            p = doc.add_paragraph(raw)
            style_docx_paragraph(p)

    model_cols = ["modelo", "accuracy", "precision", "recall", "f1", "roc_auc", "pr_auc"]
    add_docx_table(doc, data["clf"][model_cols], "Tabla 1. Benchmark de clasificacion")
    cluster_cols = ["cluster", "observaciones", "denuncias_tasa_100k", "victimization", "percepcion", "confianza", "brecha"]
    add_docx_table(doc, data["clusters"][cluster_cols], "Tabla 2. Perfil promedio de clusters")

    figure_plan = [
        ("01_evolucion_anual_denuncias.png", "Figura 1. Evolucion anual de denuncias policiales registradas."),
        ("02_tasa_denuncias_departamento_2024.png", "Figura 2. Tasa departamental de denuncias por 100,000 habitantes en 2024."),
        ("05_scatter_victimizacion_percepcion.png", "Figura 3. Relacion entre victimizacion y percepcion de inseguridad."),
        ("08_pca_clusters.png", "Figura 4. Visualizacion PCA de clusters territoriales."),
        ("15_feature_importance.png", "Figura 5. Importancia por permutacion del mejor clasificador."),
    ]
    for name, caption in figure_plan:
        path = FIGURES / name
        if path.exists():
            doc.add_paragraph(caption)
            doc.add_picture(str(path), width=Cm(15.8))

    doc.save(REPORT / "informe_tecnico.docx")


def pdf_paragraph(text: str, styles) -> Paragraph:
    return Paragraph(text.replace("\n", "<br/>"), styles["Body"])


def write_pdf() -> None:
    data = load_inputs()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Title2", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, spaceAfter=8, textColor=colors.HexColor("#1f3b4d")))
    styles.add(ParagraphStyle(name="Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=17, alignment=4, spaceAfter=8))
    story = []
    for title, body in paragraphs():
        story.append(Paragraph(title, styles["Title2"]))
        for raw in body.split("\n\n"):
            story.append(pdf_paragraph(raw, styles))
        story.append(Spacer(1, 0.08 * inch))

    model_cols = ["modelo", "accuracy", "precision", "recall", "f1"]
    model_data = [model_cols] + data["clf"][model_cols].round(3).astype(str).values.tolist()
    table = Table(model_data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3b4d")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONT", (0, 0), (-1, -1), "Helvetica", 8)]))
    story.append(table)
    story.append(PageBreak())

    for name, caption in [
        ("01_evolucion_anual_denuncias.png", "Figura 1. Evolucion anual de denuncias."),
        ("05_scatter_victimizacion_percepcion.png", "Figura 2. Victimizacion y percepcion."),
        ("08_pca_clusters.png", "Figura 3. Clusters territoriales."),
        ("15_feature_importance.png", "Figura 4. Importancia predictiva."),
    ]:
        path = FIGURES / name
        if path.exists():
            story.append(Paragraph(caption, styles["Body"]))
            story.append(Image(str(path), width=6.4 * inch, height=3.8 * inch))
            story.append(Spacer(1, 0.16 * inch))

    doc = SimpleDocTemplate(str(REPORT / "informe_tecnico.pdf"), pagesize=LETTER, rightMargin=0.8 * inch, leftMargin=0.8 * inch, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    doc.build(story)


def main() -> None:
    write_docx()
    write_pdf()
    print("Informe tecnico ampliado generado en report/informe_tecnico.docx y report/informe_tecnico.pdf")


if __name__ == "__main__":
    main()
