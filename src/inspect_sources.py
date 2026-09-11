from __future__ import annotations

import csv
import json
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
TABLES = ROOT / "outputs" / "tables"


CSV_FILES = {
    "denuncias": RAW / "denuncias_policiales_2018_2026_julio.csv",
    "enapres": RAW / "victimizacion_percepcion_confianza_pnp_2013_2024.csv",
}

DICTIONARIES = {
    "denuncias": RAW / "diccionario_denuncias_policiales_abr_2026.xlsx",
    "enapres": RAW / "diccionario_victimizacion_percepcion_confianza_pnp.xlsx",
}

METADATA_DOCS = {
    "denuncias": RAW / "metadato_denuncias_policiales_julio_2026.docx",
    "enapres": RAW / "metadato_victimizacion_percepcion_confianza_pnp.docx",
}


def sniff_csv(path: Path) -> dict:
    raw = path.read_bytes()[:100_000]
    encoding = "utf-8-sig"
    try:
        sample = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "latin1"
        sample = raw.decode(encoding)
    dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
    return {"encoding": encoding, "sep": dialect.delimiter}


def read_csv(path: Path, options: dict) -> pd.DataFrame:
    return pd.read_csv(path, sep=options["sep"], encoding=options["encoding"], low_memory=False)


def profile_dataframe(name: str, df: pd.DataFrame, options: dict) -> dict:
    profile = {
        "dataset": name,
        "path": str(CSV_FILES[name]),
        "encoding": options["encoding"],
        "separator": options["sep"],
        "rows": int(df.shape[0]),
        "columns": int(df.shape[1]),
        "duplicates": int(df.duplicated().sum()),
        "columns_list": list(df.columns),
        "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
        "nulls": {col: int(df[col].isna().sum()) for col in df.columns},
        "nunique": {col: int(df[col].nunique(dropna=True)) for col in df.columns},
        "sample_rows": df.head(5).astype(str).to_dict(orient="records"),
    }
    return profile


def read_dictionary(path: Path) -> dict:
    sheets = pd.read_excel(path, sheet_name=None)
    return {
        sheet: frame.fillna("").astype(str).to_dict(orient="records")
        for sheet, frame in sheets.items()
    }


def read_docx_text(path: Path) -> str:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))
    return "\n".join(paragraphs + table_lines)


def write_tables(profiles: dict, dictionaries: dict, metadata_texts: dict) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    dataset_rows = []
    column_rows = []
    for name, profile in profiles.items():
        dataset_rows.append(
            {
                "dataset": name,
                "rows": profile["rows"],
                "columns": profile["columns"],
                "duplicates": profile["duplicates"],
                "encoding": profile["encoding"],
                "separator": profile["separator"],
                "column_names": ", ".join(profile["columns_list"]),
            }
        )
        for col in profile["columns_list"]:
            column_rows.append(
                {
                    "dataset": name,
                    "column": col,
                    "dtype": profile["dtypes"][col],
                    "nulls": profile["nulls"][col],
                    "nunique": profile["nunique"][col],
                }
            )

    pd.DataFrame(dataset_rows).to_csv(TABLES / "fase1_dataset_profile.csv", index=False)
    pd.DataFrame(column_rows).to_csv(TABLES / "fase1_column_profile.csv", index=False)
    for name, sheets in dictionaries.items():
        rows = []
        for sheet, records in sheets.items():
            for record in records:
                record = {"sheet": sheet, **record}
                rows.append(record)
        pd.DataFrame(rows).to_csv(TABLES / f"fase1_diccionario_{name}.csv", index=False)
    for name, text in metadata_texts.items():
        (TABLES / f"fase1_metadato_{name}.txt").write_text(text, encoding="utf-8")


def main() -> None:
    profiles = {}
    dictionaries = {}
    metadata_texts = {}

    for name, path in CSV_FILES.items():
        options = sniff_csv(path)
        df = read_csv(path, options)
        profiles[name] = profile_dataframe(name, df, options)

    for name, path in DICTIONARIES.items():
        dictionaries[name] = read_dictionary(path)

    for name, path in METADATA_DOCS.items():
        metadata_texts[name] = read_docx_text(path)

    write_tables(profiles, dictionaries, metadata_texts)
    payload = {
        "profiles": profiles,
        "dictionaries": dictionaries,
        "metadata_texts": metadata_texts,
    }
    (TABLES / "fase1_source_inspection.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps({name: profiles[name] for name in profiles}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
