from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "defense" / "guia_simple_para_entender_el_proyecto.md"
OUTPUT = ROOT / "defense" / "guia_simple_para_entender_el_proyecto.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def set_run_font(run, size=10.5, bold=False, color="111111") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=10.5, bold=False, color="111111") -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.08


def add_text_with_bold(paragraph, text: str, size=10.5) -> None:
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=(idx % 2 == 1))


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def build_docx() -> None:
    markdown = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    for name, size, bold in [
        ("Normal", 10.5, False),
        ("Title", 24, True),
        ("Heading 1", 16, True),
        ("Heading 2", 13, True),
    ]:
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text_with_bold(title, "Guia simple para entender el proyecto", size=24)
    doc.add_paragraph(
        "Material de estudio para explicar el proyecto de criminalidad, seguridad ciudadana y Machine Learning sin necesidad de tener formacion en ingenieria."
    )

    idx = 0
    while idx < len(markdown):
        line = markdown[idx].rstrip()
        if not line:
            idx += 1
            continue
        if line.startswith("# "):
            idx += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_text_with_bold(paragraph, line[3:].strip(), size=16)
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_text_with_bold(paragraph, line[4:].strip(), size=13)
        elif line.startswith("| "):
            table_lines = []
            while idx < len(markdown) and markdown[idx].startswith("| "):
                table_lines.append(markdown[idx])
                idx += 1
            idx -= 1
            rows = parse_table([table_lines[0], *table_lines[2:]])
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_borders(cell)
                    if r == 0:
                        set_cell_shading(cell, "1F4E79")
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        run = paragraph.add_run(value)
                        set_run_font(run, size=9.5, bold=(r == 0), color=("FFFFFF" if r == 0 else "111111"))
            doc.add_paragraph()
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_text_with_bold(paragraph, line[2:].strip(), size=10.5)
            style_paragraph(paragraph)
        elif len(line) > 2 and line[0].isdigit() and ". " in line[:4]:
            paragraph = doc.add_paragraph(style="List Number")
            add_text_with_bold(paragraph, line.split(". ", 1)[1].strip(), size=10.5)
            style_paragraph(paragraph)
        elif line.startswith('"') and line.endswith('"'):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            add_text_with_bold(paragraph, line, size=11)
            style_paragraph(paragraph, size=11, bold=False, color="1F4E79")
        else:
            paragraph = doc.add_paragraph()
            add_text_with_bold(paragraph, line, size=10.5)
            style_paragraph(paragraph)
        idx += 1

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
